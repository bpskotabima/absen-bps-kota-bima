import json
import math
import os
import sqlite3
from datetime import datetime, time
from functools import wraps

from flask import (Flask, g, jsonify, redirect, render_template, request,
                   session, url_for)

# [TAMBAHAN] import untuk rekap & ekspor Excel
import calendar
from io import BytesIO
from flask import render_template_string, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "ganti-kunci-rahasia-ini")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin123")  # ganti!
DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "absensi.db")

# ---- Aturan jam: absen hanya 1x, pada jam istirahat ----
JAM_ISTIRAHAT = time(12, 0)   # absen dibuka pukul 12.00
BATAS_ABSEN = time(13, 0)     # lewat 13.00 = terlambat, wajib isi alasan
BATAS_WAJAH = 0.5             # makin kecil makin ketat (0.4 - 0.6 umum dipakai)

JENIS = "istirahat"
LABEL = {
    "istirahat": "Absen jam istirahat",
    # label untuk data lama (kalau ada sisa tes sebelumnya)
    "masuk": "Masuk kerja (data lama)",
    "kembali": "Masuk kembali (data lama)",
    "pulang": "Pulang (data lama)",
}


def sekarang():
    """Waktu sesuai jam komputer yang menjalankan aplikasi."""
    return datetime.now()


# ---------- Database ----------
def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def tutup_db(_):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    db = sqlite3.connect(DB_PATH)
    db.executescript(
        """
        CREATE TABLE IF NOT EXISTS karyawan (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nip TEXT NOT NULL UNIQUE,
            nama TEXT NOT NULL,
            descriptor TEXT NOT NULL,          -- 128 angka wajah (JSON)
            dibuat TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS absensi (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            karyawan_id INTEGER NOT NULL REFERENCES karyawan(id) ON DELETE CASCADE,
            tanggal TEXT NOT NULL,             -- YYYY-MM-DD
            jenis TEXT NOT NULL,
            waktu TEXT NOT NULL,               -- HH:MM:SS
            terlambat INTEGER NOT NULL DEFAULT 0,
            alasan TEXT,
            UNIQUE (karyawan_id, tanggal, jenis)
        );
        """
    )
    db.commit()
    db.close()


def admin_wajib(fn):
    @wraps(fn)
    def bungkus(*a, **kw):
        if not session.get("admin"):
            return redirect(url_for("login"))
        return fn(*a, **kw)
    return bungkus


def jarak(a, b):
    return math.sqrt(sum((x - y) ** 2 for x, y in zip(a, b)))


def valid_descriptor(d):
    return (isinstance(d, list) and len(d) == 128
            and all(isinstance(x, (int, float)) for x in d))


# ---------- Halaman ----------
@app.route("/")
def halaman_absen():
    return render_template("absen.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    galat = None
    if request.method == "POST":
        if request.form.get("password") == ADMIN_PASSWORD:
            session["admin"] = True
            return redirect(url_for("admin"))
        galat = "Password salah."
    return render_template("login.html", galat=galat)


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("halaman_absen"))


@app.route("/admin")
@admin_wajib
def admin():
    db = get_db()
    tanggal = request.args.get("tanggal") or sekarang().strftime("%Y-%m-%d")
    karyawan = db.execute(
        "SELECT id, nip, nama, dibuat FROM karyawan ORDER BY nama").fetchall()
    rekap = db.execute(
        """SELECT k.nama, k.nip, a.jenis, a.waktu, a.terlambat, a.alasan
           FROM absensi a JOIN karyawan k ON k.id = a.karyawan_id
           WHERE a.tanggal = ? ORDER BY k.nama, a.waktu""", (tanggal,)).fetchall()
    return render_template("admin.html", karyawan=karyawan, rekap=rekap,
                           tanggal=tanggal, label=LABEL)


@app.route("/admin/daftar")
@admin_wajib
def halaman_daftar():
    return render_template("daftar.html")


@app.route("/admin/hapus/<int:kid>", methods=["POST"])
@admin_wajib
def hapus_karyawan(kid):
    db = get_db()
    db.execute("PRAGMA foreign_keys = ON")
    db.execute("DELETE FROM absensi WHERE karyawan_id = ?", (kid,))
    db.execute("DELETE FROM karyawan WHERE id = ?", (kid,))
    db.commit()
    return redirect(url_for("admin"))


# ---------- API ----------
@app.route("/api/daftar", methods=["POST"])
@admin_wajib
def api_daftar():
    data = request.get_json(silent=True) or {}
    nip = (data.get("nip") or "").strip()
    nama = (data.get("nama") or "").strip()
    desc = data.get("descriptor")
    if not nip or not nama or not valid_descriptor(desc):
        return jsonify(pesan="Data tidak lengkap atau wajah tidak terbaca."), 400
    db = get_db()
    try:
        db.execute(
            "INSERT INTO karyawan (nip, nama, descriptor, dibuat) VALUES (?,?,?,?)",
            (nip, nama, json.dumps(desc), sekarang().isoformat(timespec="seconds")))
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify(pesan="NIP sudah terdaftar."), 409
    return jsonify(pesan=f"{nama} berhasil didaftarkan.")


@app.route("/api/absen", methods=["POST"])
def api_absen():
    data = request.get_json(silent=True) or {}
    desc = data.get("descriptor")
    alasan = (data.get("alasan") or "").strip()
    if not valid_descriptor(desc):
        return jsonify(status="galat", pesan="Wajah tidak terbaca. Coba lagi."), 400

    db = get_db()

    # 1) Kenali wajah: cari karyawan dengan jarak terdekat
    terbaik, jarak_min = None, 9.9
    for k in db.execute("SELECT id, nama, descriptor FROM karyawan"):
        j = jarak(desc, json.loads(k["descriptor"]))
        if j < jarak_min:
            terbaik, jarak_min = k, j
    if terbaik is None or jarak_min > BATAS_WAJAH:
        return jsonify(status="galat", pesan="Wajah tidak dikenali."), 404

    nama = terbaik["nama"]
    skrg = sekarang()
    tanggal = skrg.strftime("%Y-%m-%d")
    menit_ini = skrg.time().replace(second=0, microsecond=0)

    # 2) Tidak boleh absen dua kali dalam sehari
    sudah = db.execute(
        "SELECT waktu FROM absensi WHERE karyawan_id=? AND tanggal=? AND jenis=?",
        (terbaik["id"], tanggal, JENIS)).fetchone()
    if sudah:
        return jsonify(
            status="galat",
            pesan=f"{nama}, Anda sudah absen hari ini pukul {sudah['waktu']}. "
                  "Absen hanya boleh satu kali."), 409

    # 3) Absen baru dibuka pukul 12.00
    if menit_ini < JAM_ISTIRAHAT:
        return jsonify(status="galat",
                       pesan="Absen jam istirahat baru dibuka pukul 12.00."), 400

    # 4) Lewat pukul 13.00 = terlambat, wajib isi alasan
    terlambat = int(menit_ini > BATAS_ABSEN)
    if terlambat and not alasan:
        return jsonify(
            status="butuh_alasan", nama=nama,
            pesan=f"{nama}, Anda absen setelah pukul 13.00 (jam istirahat sudah selesai). "
                  "Tulis alasan terlambat."), 200

    # 5) Simpan
    try:
        db.execute(
            """INSERT INTO absensi (karyawan_id, tanggal, jenis, waktu, terlambat, alasan)
               VALUES (?,?,?,?,?,?)""",
            (terbaik["id"], tanggal, JENIS, skrg.strftime("%H:%M:%S"),
             terlambat, alasan if terlambat else None))
        db.commit()
    except sqlite3.IntegrityError:
        return jsonify(status="galat",
                       pesan=f"{nama}, Anda sudah absen hari ini."), 409

    return jsonify(status="ok", nama=nama, jenis=LABEL[JENIS],
                   waktu=skrg.strftime("%H:%M:%S"), terlambat=bool(terlambat),
                   pesan=f"Absen jam istirahat tercatat untuk {nama}.")


# =====================================================================
# [TAMBAHAN 1] Efek halus & modern untuk SEMUA halaman HTML.
# Disuntikkan otomatis lewat after_request, jadi template tidak perlu diubah:
#  - transisi antar-halaman yang mulus (View Transitions)
#  - konten muncul dengan fade + naik, bertahap
#  - bar loading di atas layar saat pindah halaman / kirim form
#  - efek riak (ripple) pada tombol
# =====================================================================
EFEK_HEAD = """<style id="efek-halus">
@view-transition{navigation:auto}
html{scroll-behavior:smooth}
@keyframes masuk-halus{from{opacity:0;transform:translateY(12px)}to{opacity:1;transform:none}}
main>*{animation:masuk-halus .55s cubic-bezier(.2,.8,.2,1) backwards}
main>*:nth-child(2){animation-delay:.05s}
main>*:nth-child(3){animation-delay:.10s}
main>*:nth-child(4){animation-delay:.15s}
main>*:nth-child(5){animation-delay:.20s}
main>*:nth-child(n+6){animation-delay:.25s}
#bar-muat{position:fixed;top:0;left:0;height:3px;width:0;opacity:0;z-index:9999;
  background:linear-gradient(90deg,#6366f1,#22d3ee);box-shadow:0 0 10px #6366f1;
  transition:width .6s ease,opacity .3s}
button,.tombol{position:relative;overflow:hidden}
.riak{position:absolute;border-radius:50%;background:rgba(255,255,255,.4);
  transform:scale(0);animation:riak .6s linear;pointer-events:none}
@keyframes riak{to{transform:scale(4);opacity:0}}
@media (prefers-reduced-motion:reduce){
  main>*,.riak{animation:none!important}html{scroll-behavior:auto}
}
</style>"""

EFEK_BODY = """<script id="efek-halus-js">
(function(){
  var bar=document.createElement('div');bar.id='bar-muat';document.body.appendChild(bar);
  var timer=null;
  function mulai(){
    bar.style.opacity='1';bar.style.width='72%';
    clearTimeout(timer);timer=setTimeout(selesai,8000);
  }
  function selesai(){bar.style.width='100%';setTimeout(function(){bar.style.opacity='0';bar.style.width='0';},300);}
  window.addEventListener('pageshow',selesai);

  document.addEventListener('submit',function(e){
    var url=(e.submitter&&e.submitter.formAction)||e.target.action||'';
    if(e.defaultPrevented||url.indexOf('/ekspor')>-1)return;   // unduhan file tidak pindah halaman
    mulai();
  });

  document.addEventListener('click',function(e){
    var b=e.target.closest('button,.tombol');
    if(b&&!b.disabled){
      var r=b.getBoundingClientRect(),s=Math.max(r.width,r.height),o=document.createElement('span');
      o.className='riak';o.style.width=o.style.height=s+'px';
      o.style.left=(e.clientX-r.left-s/2)+'px';o.style.top=(e.clientY-r.top-s/2)+'px';
      b.appendChild(o);setTimeout(function(){o.remove();},650);
    }
    var a=e.target.closest('a[href]');
    if(a&&a.origin===location.origin&&!a.target&&!a.hasAttribute('download')
       &&a.pathname.indexOf('/ekspor')<0&&a.getAttribute('href').charAt(0)!=='#'
       &&!e.ctrlKey&&!e.metaKey&&!e.shiftKey&&e.button===0){mulai();}
  });
})();
</script>"""


@app.after_request
def efek_halus(resp):
    if (resp.status_code == 200 and resp.mimetype == "text/html"
            and not resp.direct_passthrough):
        html = resp.get_data(as_text=True)
        if "</head>" in html and "</body>" in html and "efek-halus" not in html:
            html = html.replace("</head>", EFEK_HEAD + "</head>", 1)
            html = html.replace("</body>", EFEK_BODY + "</body>", 1)
            resp.set_data(html)
    return resp


# =====================================================================
# [TAMBAHAN 2] Halaman rekap bulanan (tampilan modern) + ekspor Excel
# =====================================================================
NAMA_BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli",
              "Agustus", "September", "Oktober", "November", "Desember"]
NAMA_HARI = ["Sen", "Sel", "Rab", "Kam", "Jum", "Sab", "Min"]
NAMA_HARI_PANJANG = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]

HALAMAN_REKAP = """<!doctype html>
<html lang="id">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Rekap Absensi {{ nama_bulan }} {{ tahun }}</title>
<link rel="stylesheet" href="{{ url_for('static', filename='style.css') }}">
<style>
  .ringkasan{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:12px;margin:18px 0}
  .ringkasan .kartu{padding:14px 16px;position:relative;overflow:hidden}
  .ringkasan .kartu::before{content:"";position:absolute;left:0;top:0;bottom:0;width:4px;background:var(--utama)}
  .ringkasan .ok::before{background:var(--ok)}
  .ringkasan .tl::before{background:var(--bahaya)}
  .ringkasan .tt::before{background:var(--peringatan)}
  .ringkasan small{color:var(--redup);font-weight:600;font-size:.78rem;text-transform:uppercase;letter-spacing:.05em}
  .ringkasan .angka{font-size:1.9rem;font-weight:800;letter-spacing:-.02em;line-height:1.2}
  .filter{display:flex;gap:10px;align-items:center;flex-wrap:wrap}
  .filter label{margin:0}
  .filter input{width:auto}
  .rekap th,.rekap td{padding:7px 8px;text-align:center;white-space:nowrap;font-size:.82rem}
  .rekap td.kiri,.rekap th.kiri{text-align:left}
  table.tabel.rekap tr td.libur{background:var(--garis);opacity:.75}
  table.tabel.rekap tr td.telat{background:var(--bahaya-muda)}
  table.tabel.rekap tr td.ringkas{font-weight:700;background:var(--utama-muda)}
  .ket{margin-top:14px;font-size:.85rem;color:var(--redup);line-height:1.7}
</style>
</head>
<body>
<main class="lebar">
  <div class="bar"><h1>Rekap absensi</h1>
    <nav><a href="{{ url_for('admin') }}">&larr; Admin</a>
      <a href="{{ url_for('halaman_absen') }}">Halaman absen</a>
      <a href="{{ url_for('logout') }}">Keluar</a></nav></div>
  <p class="redup">Bulan {{ nama_bulan }} {{ tahun }} &middot; {{ baris|length }} karyawan</p>

  <div class="kartu">
    <form method="get" action="{{ url_for('halaman_rekap') }}" class="filter">
      <label for="bulan">Bulan</label>
      <input type="month" id="bulan" name="bulan" value="{{ bulan }}" required>
      <button type="submit" class="sekunder">Tampilkan</button>
      <a class="tombol" href="{{ url_for('ekspor_excel', bulan=bulan) }}">&#11015; Unduh Excel (.xlsx)</a>
    </form>
  </div>

  <div class="ringkasan">
    <div class="kartu"><small>Total hadir</small><div class="angka">{{ total.hadir }}</div></div>
    <div class="kartu ok"><small>Tepat waktu</small><div class="angka">{{ total.tepat }}</div></div>
    <div class="kartu tl"><small>Terlambat</small><div class="angka">{{ total.telat }}</div></div>
    <div class="kartu tt"><small>Tidak absen</small><div class="angka">{{ total.tidak }}</div></div>
  </div>

  <div class="gulir">
  <table class="tabel rekap">
    <tr>
      <th>No</th><th class="kiri">NIP</th><th class="kiri">Nama</th>
      {% for h in hari %}<th>{{ h.d }}<br>{{ h.n }}</th>{% endfor %}
      <th>Hadir</th><th>Tepat</th><th>Telat</th><th>Alpa</th>
    </tr>
    {% for b in baris %}
    <tr>
      <td>{{ loop.index }}</td>
      <td class="kiri">{{ b.nip }}</td>
      <td class="kiri">{{ b.nama }}</td>
      {% for s in b.sel %}<td class="{{ s.kelas }}">{{ s.teks }}</td>{% endfor %}
      <td class="ringkas">{{ b.hadir }}</td>
      <td class="ringkas">{{ b.tepat }}</td>
      <td class="ringkas">{{ b.telat }}</td>
      <td class="ringkas">{{ b.tidak }}</td>
    </tr>
    {% else %}
    <tr><td colspan="{{ 7 + hari|length }}">Belum ada karyawan terdaftar.</td></tr>
    {% endfor %}
  </table></div>

  <p class="ket">
    Angka jam (misal 12:15) = jam absen. Merah = terlambat (lewat 13.00). Tanda &ldquo;-&rdquo; = tidak absen
    pada hari kerja. Abu-abu = Sabtu/Minggu. Hari libur nasional tidak dikenali otomatis dan dihitung
    sebagai hari kerja.
  </p>
</main>
</body>
</html>
"""


def _hitung_rekap(bulan):
    """Kumpulkan data rekap satu bulan (dipakai halaman & ekspor)."""
    tgl_awal = datetime.strptime(bulan, "%Y-%m")
    jml_hari = calendar.monthrange(tgl_awal.year, tgl_awal.month)[1]
    db = get_db()
    karyawan = db.execute("SELECT id, nip, nama FROM karyawan ORDER BY nama").fetchall()
    data = db.execute(
        """SELECT karyawan_id, tanggal, waktu, terlambat, alasan
           FROM absensi WHERE tanggal BETWEEN ? AND ? AND jenis = ?
           ORDER BY tanggal""",
        (f"{bulan}-01", f"{bulan}-{jml_hari:02d}", JENIS)).fetchall()
    absen = {(r["karyawan_id"], int(r["tanggal"][8:10])): r for r in data}
    return tgl_awal, jml_hari, karyawan, absen


def _tidak_absen(tgl, skrg):
    """Hari kerja yang sudah lewat tanpa absen."""
    return tgl.weekday() < 5 and (
        tgl.date() < skrg.date()
        or (tgl.date() == skrg.date() and skrg.time() > BATAS_ABSEN))


@app.route("/admin/rekap")
@admin_wajib
def halaman_rekap():
    bulan = request.args.get("bulan") or sekarang().strftime("%Y-%m")
    try:
        tgl_awal, jml_hari, karyawan, absen = _hitung_rekap(bulan)
    except ValueError:
        return "Format bulan tidak valid.", 400
    skrg = sekarang()

    hari = [{"d": d, "n": NAMA_HARI[tgl_awal.replace(day=d).weekday()]}
            for d in range(1, jml_hari + 1)]
    baris = []
    total = {"hadir": 0, "tepat": 0, "telat": 0, "tidak": 0}
    for k in karyawan:
        sel, hadir, tepat, telat, tidak = [], 0, 0, 0, 0
        for d in range(1, jml_hari + 1):
            tgl = tgl_awal.replace(day=d)
            rec = absen.get((k["id"], d))
            teks, kelas = "", ""
            if rec:
                hadir += 1
                teks = rec["waktu"][:5]
                if rec["terlambat"]:
                    telat += 1
                    kelas = "telat"
                else:
                    tepat += 1
            elif _tidak_absen(tgl, skrg):
                tidak += 1
                teks = "-"
            if tgl.weekday() >= 5 and not rec:
                kelas = "libur"
            sel.append({"teks": teks, "kelas": kelas})
        baris.append({"nip": k["nip"], "nama": k["nama"], "sel": sel,
                      "hadir": hadir, "tepat": tepat, "telat": telat, "tidak": tidak})
        total["hadir"] += hadir
        total["tepat"] += tepat
        total["telat"] += telat
        total["tidak"] += tidak

    return render_template_string(
        HALAMAN_REKAP, bulan=bulan, tahun=tgl_awal.year,
        nama_bulan=NAMA_BULAN[tgl_awal.month - 1], hari=hari, baris=baris, total=total)


@app.route("/admin/ekspor")
@admin_wajib
def ekspor_excel():
    bulan = request.args.get("bulan") or sekarang().strftime("%Y-%m")
    try:
        tgl_awal, jml_hari, karyawan, absen = _hitung_rekap(bulan)
    except ValueError:
        return "Format bulan tidak valid.", 400
    tahun, bln = tgl_awal.year, tgl_awal.month
    skrg = sekarang()
    db = get_db()

    # --- gaya ---
    tipis = Side(style="thin", color="B7BEC8")
    garis = Border(left=tipis, right=tipis, top=tipis, bottom=tipis)
    tengah = Alignment(horizontal="center", vertical="center", wrap_text=True)
    kiri = Alignment(horizontal="left", vertical="center", wrap_text=True)
    isi_header = PatternFill("solid", fgColor="14639E")
    font_header = Font(bold=True, color="FFFFFF")
    isi_akhir_pekan = PatternFill("solid", fgColor="E3E6EA")
    isi_telat = PatternFill("solid", fgColor="FDE2E0")
    font_telat = Font(bold=True, color="B3261E")
    isi_ringkas = PatternFill("solid", fgColor="EEF1F5")

    wb = Workbook()

    # ========== Sheet 1: Rekap per karyawan ==========
    ws = wb.active
    ws.title = "Rekap Bulanan"
    kolom_ringkas = ["Hadir", "Tepat Waktu", "Terlambat", "Tidak Absen"]
    total_kolom = 3 + jml_hari + len(kolom_ringkas)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_kolom)
    ws["A1"] = "REKAP ABSENSI JAM ISTIRAHAT"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_kolom)
    ws["A2"] = f"Bulan {NAMA_BULAN[bln - 1]} {tahun}"
    ws["A2"].font = Font(bold=True, size=12)
    ws["A2"].alignment = Alignment(horizontal="center")

    baris_header = 4
    for c, h in enumerate(["No", "NIP", "Nama"], start=1):
        ws.cell(baris_header, c, h)
    for d in range(1, jml_hari + 1):
        hari = tgl_awal.replace(day=d).weekday()
        ws.cell(baris_header, 3 + d, f"{d}\n{NAMA_HARI[hari]}")
    for i, h in enumerate(kolom_ringkas):
        ws.cell(baris_header, 4 + jml_hari + i, h)
    for c in range(1, total_kolom + 1):
        sel = ws.cell(baris_header, c)
        sel.fill, sel.font, sel.alignment, sel.border = isi_header, font_header, tengah, garis
    ws.row_dimensions[baris_header].height = 32

    for n, k in enumerate(karyawan, start=1):
        r = baris_header + n
        ws.cell(r, 1, n).alignment = tengah
        ws.cell(r, 2, k["nip"]).alignment = kiri
        ws.cell(r, 3, k["nama"]).alignment = kiri
        hadir = tepat = telat = tidak = 0
        for d in range(1, jml_hari + 1):
            tgl = tgl_awal.replace(day=d)
            sel = ws.cell(r, 3 + d)
            sel.alignment = tengah
            rec = absen.get((k["id"], d))
            if rec:
                hadir += 1
                sel.value = rec["waktu"][:5]
                if rec["terlambat"]:
                    telat += 1
                    sel.fill, sel.font = isi_telat, font_telat
                else:
                    tepat += 1
            elif _tidak_absen(tgl, skrg):
                tidak += 1
                sel.value = "-"
            if tgl.weekday() >= 5 and not rec:
                sel.fill = isi_akhir_pekan
        for i, nilai in enumerate([hadir, tepat, telat, tidak]):
            sel = ws.cell(r, 4 + jml_hari + i, nilai)
            sel.alignment, sel.fill = tengah, isi_ringkas
            sel.font = Font(bold=True)
        for c in range(1, total_kolom + 1):
            ws.cell(r, c).border = garis

    if not karyawan:
        ws.cell(baris_header + 1, 1, "Belum ada karyawan terdaftar.")

    r_ket = baris_header + len(karyawan) + 2
    ws.cell(r_ket, 1, "Keterangan:").font = Font(bold=True)
    ws.cell(r_ket + 1, 1, "Angka jam (misal 12:15) = jam absen. Merah = terlambat (lewat 13.00).")
    ws.cell(r_ket + 2, 1, "Tanda “-” = tidak absen pada hari kerja. Abu-abu = Sabtu/Minggu. "
                          "Hari libur nasional tidak dikenali otomatis dan dihitung sebagai hari kerja.")

    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 28
    for d in range(1, jml_hari + 1):
        ws.column_dimensions[get_column_letter(3 + d)].width = 6.5
    for i in range(len(kolom_ringkas)):
        ws.column_dimensions[get_column_letter(4 + jml_hari + i)].width = 12
    ws.freeze_panes = ws.cell(baris_header + 1, 4)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    # ========== Sheet 2: Detail & alasan terlambat ==========
    ws2 = wb.create_sheet("Detail Absensi")
    ws2.merge_cells("A1:H1")
    ws2["A1"] = f"DETAIL ABSENSI — {NAMA_BULAN[bln - 1].upper()} {tahun}"
    ws2["A1"].font = Font(bold=True, size=13)
    kepala = ["No", "Tanggal", "Hari", "NIP", "Nama", "Jam Absen", "Status", "Alasan Terlambat"]
    for c, h in enumerate(kepala, start=1):
        sel = ws2.cell(3, c, h)
        sel.fill, sel.font, sel.alignment, sel.border = isi_header, font_header, tengah, garis
    detail = db.execute(
        """SELECT a.tanggal, a.waktu, a.terlambat, a.alasan, k.nip, k.nama
           FROM absensi a JOIN karyawan k ON k.id = a.karyawan_id
           WHERE a.tanggal BETWEEN ? AND ? AND a.jenis = ?
           ORDER BY a.tanggal, k.nama""",
        (f"{bulan}-01", f"{bulan}-{jml_hari:02d}", JENIS)).fetchall()
    for n, r in enumerate(detail, start=1):
        tgl = datetime.strptime(r["tanggal"], "%Y-%m-%d")
        nilai = [n, tgl.strftime("%d-%m-%Y"), NAMA_HARI_PANJANG[tgl.weekday()], r["nip"],
                 r["nama"], r["waktu"], "Terlambat" if r["terlambat"] else "Tepat waktu",
                 r["alasan"] or ""]
        for c, v in enumerate(nilai, start=1):
            sel = ws2.cell(3 + n, c, v)
            sel.border = garis
            sel.alignment = kiri if c in (4, 5, 8) else tengah
            if r["terlambat"] and c == 7:
                sel.font, sel.fill = font_telat, isi_telat
    if not detail:
        ws2.cell(4, 1, "Belum ada data absensi pada bulan ini.")
    for huruf, lebar in zip("ABCDEFGH", [5, 13, 11, 18, 28, 11, 14, 45]):
        ws2.column_dimensions[huruf].width = lebar
    ws2.freeze_panes = "A4"
    ws2.page_setup.orientation = "landscape"
    ws2.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws2.page_setup.fitToWidth = 1
    ws2.page_setup.fitToHeight = 0

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True,
                     download_name=f"rekap-absensi-{bulan}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# Pastikan tabel database ada (juga bila dijalankan dengan `flask run`)
init_db()


if __name__ == "__main__":
    init_db()
    app.run(host="0.0.0.0", port=5000, debug=True)